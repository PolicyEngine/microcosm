#!/usr/bin/env python
"""Build uk/centax_childcare_parent_facts.json from CenTax's released workbook.

Advani, Flew, Pepin-Hall and Summers (CenTax, September 2026), *Removing the
childcare cliff-edge*, released their HMRC tabulations as
``AdvaniFlewPepinhallSummers2026_ChildcareThreshold_data.xlsx``; its SHA-256 is
pinned in ``uk_runtime/centax_childcare_parents.py``. This tool reads
Figures 3.1, 3.3, 3.6 and 3.7 from it and writes nine derived facts (see that
module): four ANI band shares of the parents in (0, £200k], and five splits of
the parents in (£80k, £150k], each times ``parent_count``.

``parent_count`` is the baseline build's own number of such parents at the
calibration year, read from the JSON ``tools/diagnose_uk_centax_parents.py``
writes on the baseline H5; the resource records that H5's SHA-256. The
uprating that ages the 2021-22 edges is the installed policyengine-uk's OBR
average earnings index, which must agree with the diagnostics'.

Kept as diagnostics, never targets: Table A.1's levels (1,648,632 parents; 77,615
on £80k-£140k), Figure 2.1's live childcare account shares and Figure 3.12's
non-working-partner shares. CenTax's model outputs (Tables 3.1, 3.2 and 4.1,
the frictions, elasticities, marginal bunchers and every cost) are not read.

    uv run --no-sync python tools/build_uk_centax_childcare_parent_facts.py \\
        --workbook AdvaniFlewPepinhallSummers2026_ChildcareThreshold_data.xlsx \\
        --baseline-diagnostics centax_parents_a6299504_2025.json
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections.abc import Mapping, Sequence
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

from microcosm.build.uk_runtime.centax_childcare_parents import (
    CENTAX_ANI_BANDS_2022,
    CENTAX_DATA_TAX_YEAR,
    CENTAX_DERIVED_FACT_KEY_PREFIX,
    CENTAX_FACTS_RESOURCE,
    CENTAX_REGION_GROUPS,
    CENTAX_REPORT_URL,
    CENTAX_SPLIT_BAND_2022,
    CENTAX_SPLITS,
    CENTAX_WORKBOOK_SHA256,
    RESOURCE_KIND,
    ROW_SCHEMA_VERSION,
    centax_ani_uprating,
    centax_cell_ids,
    load_uk_centax_childcare_parent_facts,
    model_ani_edges,
    rows_sha256,
)

GENERATOR = {
    "tool": "tools/build_uk_centax_childcare_parent_facts.py",
    "version": "uk-centax-childcare-parents-v1",
}
DEFAULT_OUT = (
    Path(__file__).resolve().parents[1]
    / "packages/microcosm-build/src/microcosm/build/uk"
    / CENTAX_FACTS_RESOURCE
)
WORKBOOK_NAME = "AdvaniFlewPepinhallSummers2026_ChildcareThreshold_data.xlsx"
#: Sheet, expected header row (from column B) and the bin step in pounds.
SHEETS: Mapping[str, tuple[tuple[str, ...], int]] = {
    "Fig. 3.1": (("Income Bin (ANI)", "Number of Taxpayer in Income Bin"), 500),
    "Fig. 3.3": (
        (
            "Income Bin (ANI)",
            "Number of Taxpayers in Income Bin (Employees)",
            "Number of Taxpayers in Income Bin (Non-Employees)",
        ),
        1000,
    ),
    "Fig. 3.6": (
        (
            "Income Bin (ANI)",
            "Number of Taxpayers in Income Bin (Males)",
            "Number of Taxpayers in Income Bin (Females)",
        ),
        500,
    ),
    "Fig. 3.7": (
        (
            "Income Bin (ANI)",
            "Number of Taxpayers in Income Bin (London)",
            "Number of Taxpayers in Income Bin (North and Midlands)",
            "Number of Taxpayers in Income Bin (South excl. London)",
            "Number of Taxpayers in Income Bin (Scotland)",
        ),
        1000,
    ),
}
_FIG_3_7_COLUMN = {
    "london": "Number of Taxpayers in Income Bin (London)",
    "north_and_midlands": "Number of Taxpayers in Income Bin (North and Midlands)",
    "south_excl_london": "Number of Taxpayers in Income Bin (South excl. London)",
}
_CELL_LABELS = {
    "ani_0_to_50k": "ANI up to £50,000",
    "ani_50k_to_80k": "ANI over £50,000 to £80,000",
    "ani_80k_to_120k": "ANI over £80,000 to £120,000",
    "ani_120k_to_200k": "ANI over £120,000 to £200,000",
    "female_ani_80k_to_150k": "women, ANI over £80,000 to £150,000",
    "non_employee_ani_80k_to_150k": (
        "main income source other than employment, ANI over £80,000 to £150,000"
    ),
    "london_ani_80k_to_150k": "London, ANI over £80,000 to £150,000",
    "south_excl_london_ani_80k_to_150k": (
        "South excluding London, ANI over £80,000 to £150,000"
    ),
    "north_and_midlands_ani_80k_to_150k": (
        "North and Midlands, ANI over £80,000 to £150,000"
    ),
}
_CELL_FIGURES = {
    **{band: "Fig. 3.1" for band, _, _ in CENTAX_ANI_BANDS_2022},
    **{f"{split}_ani_80k_to_150k": figure for split, figure in CENTAX_SPLITS},
}
SOURCE_MEASURE_ID = "parents_of_eligible_3_4_year_olds"
SOURCE_CONCEPT = "microcosm.centax_childcare_parents.count"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_series(workbook: Any, sheet: str) -> dict[str, dict[int, float]]:
    """A figure's columns by header, each a map from bin upper edge to count.

    Refuses a sheet whose header differs from the one this tool was written
    against, a non-numeric cell inside the bins, or bins that are not one
    step apart from the first step to the last.
    """

    header, step = SHEETS[sheet]
    rows = list(workbook[sheet].iter_rows(values_only=True))
    try:
        start = next(
            index
            for index, row in enumerate(rows)
            if tuple(row[1 : 1 + len(header)]) == header
        )
    except StopIteration:
        raise ValueError(f"{sheet}: header {header!r} not found.") from None
    series: dict[str, dict[int, float]] = {column: {} for column in header[1:]}
    labels: list[int] = []
    for row in rows[start + 1 :]:
        label = row[1]
        if label is None:
            continue
        if not isinstance(label, (int, float)) or isinstance(label, bool):
            raise ValueError(f"{sheet}: non-numeric bin label {label!r}.")
        labels.append(int(label))
        for offset, column in enumerate(header[1:], start=2):
            value = row[offset]
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise ValueError(f"{sheet}: bin {label} {column} is {value!r}.")
            series[column][int(label)] = float(value)
    if labels != list(range(step, labels[-1] + step, step)):
        raise ValueError(f"{sheet}: bins are not every £{step} from £{step}.")
    return series


def band_sum(series: Mapping[int, float], lower: float, upper: float) -> float:
    """Sum of the bins whose upper edge is in (lower, upper]."""

    return float(
        sum(value for label, value in series.items() if lower < label <= upper)
    )


def centax_shares(workbook: Any) -> tuple[dict[str, float], dict[str, Any]]:
    """Each cell's share of the parents in (0, £200k], and the workbook sums."""

    fig31 = read_series(workbook, "Fig. 3.1")["Number of Taxpayer in Income Bin"]
    fig33 = read_series(workbook, "Fig. 3.3")
    fig36 = read_series(workbook, "Fig. 3.6")
    fig37 = read_series(workbook, "Fig. 3.7")
    top = CENTAX_ANI_BANDS_2022[-1][2]
    if max(fig31) != top:
        raise ValueError(f"Fig. 3.1 ends at £{max(fig31)}, not £{top:,.0f}.")
    total = band_sum(fig31, 0.0, top)
    shares = {
        band: band_sum(fig31, lower, upper) / total
        for band, lower, upper in CENTAX_ANI_BANDS_2022
    }
    lower, upper = CENTAX_SPLIT_BAND_2022
    split_share = band_sum(fig31, lower, upper) / total

    def within(columns: Mapping[str, Mapping[int, float]], numerator: str) -> float:
        whole = sum(band_sum(column, lower, upper) for column in columns.values())
        return band_sum(columns[numerator], lower, upper) / whole

    england_regions = {
        group: fig37[column] for group, column in _FIG_3_7_COLUMN.items()
    }
    if set(england_regions) != set(CENTAX_REGION_GROUPS):
        raise ValueError("Fig. 3.7 groups differ from the module's region groups.")
    split_within = {
        "female": within(fig36, "Number of Taxpayers in Income Bin (Females)"),
        "non_employee": within(
            fig33, "Number of Taxpayers in Income Bin (Non-Employees)"
        ),
        **{group: within(england_regions, group) for group in england_regions},
    }
    for split, _ in CENTAX_SPLITS:
        shares[f"{split}_ani_80k_to_150k"] = split_share * split_within[split]
    sums = {
        "fig_3_1_parents_0_to_200k": total,
        "fig_3_1_parents_80k_to_150k": band_sum(fig31, lower, upper),
        "split_shares_within_80k_to_150k": split_within,
        "fig_3_6_parents_80k_to_150k": sum(
            band_sum(column, lower, upper) for column in fig36.values()
        ),
        "fig_3_3_parents_80k_to_150k": sum(
            band_sum(column, lower, upper) for column in fig33.values()
        ),
        "fig_3_7_england_parents_80k_to_150k_pooled_2019_2022": sum(
            band_sum(column, lower, upper) for column in england_regions.values()
        ),
    }
    return shares, {"split_band_share": split_share, "workbook_sums": sums}


def diagnostic_comparators(workbook: Any) -> dict[str, Any]:
    """The CenTax figures kept as diagnostics: never targets."""

    table = list(workbook["Tab. A.1"].iter_rows(values_only=True))
    parents = next(
        row
        for row in table
        if row[1] == "3- and 4-year-olds" and row[2] == "Number of Parents"
    )
    female = next(
        row
        for row in table
        if row[2] == "Share Female (%)" and table.index(row) < table.index(parents) + 5
    )
    fig312 = {
        row[1]: row[2:5]
        for row in workbook["Fig. 3.12"].iter_rows(values_only=True)
        if isinstance(row[1], (int, float)) and not isinstance(row[1], bool)
    }
    lower, upper = CENTAX_SPLIT_BAND_2022
    in_band = [values for label, values in fig312.items() if lower < label <= upper]
    return {
        "table_a1_3_4_year_olds": {
            "parents_whole_population": parents[3],
            "parents_80k_to_140k": parents[4],
            "share_female_whole_population_pct": female[3],
            "share_female_80k_to_140k_pct": female[4],
            "note": (
                "Levels of identified parents in HMRC's data for 2021-22. Not bound: "
                "HMRC identifies parents through its own linkage, so the level is not "
                "the FRS's family count, and the DfE caseloads already set the level."
            ),
        },
        "fig_3_12_non_working_partner_share_80k_to_150k_unweighted_bin_mean": {
            "men": sum(values[1] for values in in_band) / len(in_band),
            "women": sum(values[2] for values in in_band) / len(in_band),
            "note": (
                "Mean of the £2,500 bins' shares (the figure gives shares, not counts). "
                "Not bound: the denominator is parents with an identified partner, "
                "and the workbook's 'Men & Women' column repeats 'Men'."
            ),
        },
        "fig_2_1": (
            "Live DfE childcare account shares by ANI. Not bound: the frame has no "
            "childcare-account measure; the dashboard's take-up assumption is the "
            "place to compare it."
        ),
    }


def fact_row(
    *,
    cell: str,
    value: float,
    period: str,
    edges: Sequence[int],
    inputs_sha256: str,
) -> dict[str, Any]:
    record_set_spec_id = f"uk.centax_childcare_parents.{cell}.v1"
    identity = {
        "record_set_spec_id": record_set_spec_id,
        "period": period,
        "inputs_sha256": inputs_sha256,
        "method": GENERATOR["version"],
    }
    digest = hashlib.sha256(
        json.dumps(identity, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()[:24]
    label = (
        "Parents of age-eligible 3- and 4-year-olds, "
        f"{_CELL_LABELS[cell]} (2021-22 pounds), England"
    )
    return {
        "schema_version": ROW_SCHEMA_VERSION,
        "aggregate_fact_key": CENTAX_DERIVED_FACT_KEY_PREFIX + digest,
        "semantic_fact_key": (
            "microcosm.semantic_fact.centax_childcare_parents.v1:" + digest
        ),
        "value": float(round(value, 3)),
        "value_type": "float",
        "assertion": "observation",
        "provenance_class": "derived_administrative",
        "period": {"type": "calendar_year", "value": period},
        "period_coverage": {
            "start_date": f"{period}-01-01",
            "end_date": f"{period}-12-31",
            "basis": "flow",
            "notes": (
                f"CenTax's share for tax year {CENTAX_DATA_TAX_YEAR} times the "
                f"baseline build's parents at {period}; the ANI edges are aged to "
                f"{period} pounds ({edges[0]:,} to {edges[1]:,})."
            ),
        },
        "geography": {
            "level": "country",
            "id": "E92000001",
            "vintage": "current",
            "name": "England",
        },
        "entity": {"name": "person", "role": "parent"},
        "aggregation": {"method": "sum"},
        "observed_measure": {
            "source_concept": SOURCE_CONCEPT,
            "source_measure_id": SOURCE_MEASURE_ID,
            "unit": "count",
        },
        "measure": {"unit": "count"},
        "concept_alignment": {
            "source_concept": SOURCE_CONCEPT,
            "canonical_concept": "uk.person.count",
            "relation": "derived",
        },
        "dimensions": {"microcosm.centax_cell": cell},
        "dimension_labels": {
            "microcosm.centax_cell": "CenTax parent cell (2021-22 ANI pounds)"
        },
        "dimension_value_labels": {"microcosm.centax_cell": {cell: _CELL_LABELS[cell]}},
        "source": {
            "source_name": "microcosm",
            "source_table": (
                f"CenTax, Removing the childcare cliff-edge (2026), {_CELL_FIGURES[cell]}; "
                "Microcosm derivation"
            ),
            "source_file": CENTAX_FACTS_RESOURCE,
            "vintage": "centax_2026_09",
            "source_sha256": inputs_sha256,
            "url": CENTAX_REPORT_URL,
        },
        "lineage": {"source_record_id": f"microcosm.centax_childcare_parents.{cell}"},
        "layout": {
            "record_set_id": f"uk.centax_childcare_parents.{cell}",
            "record_set_spec_id": record_set_spec_id,
            "measure_id": SOURCE_MEASURE_ID,
        },
        "label": label,
    }


def build_resource(
    *,
    workbook_path: Path,
    baseline: Mapping[str, Any],
    calibration_year: int,
    uprating: float,
    built_on: str,
) -> dict[str, Any]:
    digest = _sha256(workbook_path)
    if digest != CENTAX_WORKBOOK_SHA256:
        raise ValueError(
            f"{workbook_path} is not the pinned CenTax workbook ({digest})."
        )
    edges = model_ani_edges(uprating)
    if int(baseline["year"]) != calibration_year:
        raise ValueError(
            f"the baseline diagnostics are at {baseline['year']}, not {calibration_year}."
        )
    if baseline["model_ani_edges"] != edges:
        raise ValueError("the baseline diagnostics were banded on other ANI edges.")
    if abs(float(baseline["ani_uprating"]) - uprating) > 1e-12:
        raise ValueError("the baseline diagnostics used another ANI uprating.")
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    shares, receipts = centax_shares(workbook)
    comparators = diagnostic_comparators(workbook)
    parent_count = float(baseline["parent_count"])
    inputs = {
        "workbook": {
            "filename": WORKBOOK_NAME,
            "sha256": digest,
            "report_url": CENTAX_REPORT_URL,
            "figures": ["Fig. 3.1", "Fig. 3.3", "Fig. 3.6", "Fig. 3.7"],
            "received": "Arun Advani to Vahid Ahmadi, 21 October 2026",
        },
        "baseline": {
            "dataset": baseline["dataset"],
            "sha256": baseline["dataset_sha256"],
            "diagnostics_schema": baseline["schema"],
            "policyengine_uk_version": baseline["resolver"]["policyengine_uk_version"],
        },
        "method": GENERATOR["version"],
    }
    inputs_sha256 = hashlib.sha256(
        json.dumps(inputs, sort_keys=True).encode()
    ).hexdigest()
    period = str(calibration_year)
    rows = [
        fact_row(
            cell=cell,
            value=shares[cell] * parent_count,
            period=period,
            edges=edges[cell],
            inputs_sha256=inputs_sha256,
        )
        for cell in centax_cell_ids()
    ]
    rows.sort(key=lambda row: row["aggregate_fact_key"])
    return {
        "kind": RESOURCE_KIND,
        "schema_version": 1,
        "resource": CENTAX_FACTS_RESOURCE,
        "description": (
            "Microcosm-derived counts of parents of age-eligible 3- and 4-year-olds in "
            "England by adjusted net income, from CenTax's HMRC tabulations for "
            f"{CENTAX_DATA_TAX_YEAR} (Advani, Flew, Pepin-Hall and Summers 2026, "
            "Figures 3.1, 3.3, 3.6 and 3.7): each published share times the baseline "
            "build's own parent count. Not Chronicle facts, and no CenTax model output."
        ),
        "built_on": built_on,
        "generator": GENERATOR,
        "licence": {
            "attribution": [
                "Advani, A., Flew, D., Pepin-Hall, J. and Summers, A. (2026), Removing "
                "the childcare cliff-edge, CenTax. HMRC administrative data; the "
                "authors' calculations."
            ],
        },
        "inputs": inputs,
        "inputs_sha256": inputs_sha256,
        "parameters": {
            "calibration_year": calibration_year,
            "ani_uprating": uprating,
            "ani_uprating_basis": (
                "policyengine-uk gov.economic_assumptions.indices.obr.average_earnings, "
                f"{calibration_year} over 2021"
            ),
            "ani_bands_2022": {
                band: [lower, upper] for band, lower, upper in CENTAX_ANI_BANDS_2022
            },
            "split_band_2022": list(CENTAX_SPLIT_BAND_2022),
            "model_ani_edges": edges,
            "parent_count": parent_count,
            "parent_count_basis": (
                "baseline parents of age-eligible 3- and 4-year-olds with annual ANI in "
                f"(0, {edges[CENTAX_ANI_BANDS_2022[-1][0]][1]:,}] at {calibration_year}"
            ),
            "shares": shares,
            "split_band_share": receipts["split_band_share"],
            "workbook_sums": receipts["workbook_sums"],
            "diagnostic_comparators": comparators,
            "never_targets": [
                "Tables 3.1, 3.2 and 4.1 (bunching estimates)",
                "friction parameters, elasticities and marginal bunchers",
                "every fiscal cost",
            ],
        },
        "row_count": len(rows),
        "rows_sha256": rows_sha256(rows),
        "rows": rows,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--baseline-diagnostics", type=Path, required=True)
    parser.add_argument("--calibration-year", type=int, default=2025)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--built-on", default=date.today().isoformat())
    args = parser.parse_args(argv)
    baseline = json.loads(args.baseline_diagnostics.read_text(encoding="utf-8"))
    resource = build_resource(
        workbook_path=args.workbook,
        baseline=baseline,
        calibration_year=args.calibration_year,
        uprating=centax_ani_uprating(args.calibration_year),
        built_on=args.built_on,
    )
    rendered = json.dumps(resource, indent=2, sort_keys=True, ensure_ascii=True) + "\n"
    args.out.write_text(rendered, encoding="utf-8")
    facts = load_uk_centax_childcare_parent_facts(args.out)
    print(f"wrote {args.out} ({len(facts.rows)} rows, {facts.rows_sha256[:12]})")
    for cell in centax_cell_ids():
        print(
            f"{cell:38s} share {facts.shares[cell]:.5f} value {facts.value(cell):>11,.0f}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
